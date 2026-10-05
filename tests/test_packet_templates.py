from pathlib import Path
from types import SimpleNamespace as NS
from tempfile import TemporaryDirectory
import unittest
from docxtpl import DocxTemplate
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
T=ROOT/'docassemble/ApplicationForNameChange/data/templates'
class Person:
 def __init__(self,name):
  self.name=NS(full=lambda:name)
  self.birthdate=NS(format=lambda:'January 2, 1970')
  self.address=NS(address='123 Example Street',county='Ramsey',line_two=lambda:'St. Paul, MN 55101',on_one_line=lambda:'123 Example Street, St. Paul, MN 55101')
  self.phone_number='6125550100';self.email='robin@example.com'
  self.release_other_names='Old Alias';self.release_sex='Female';self.release_race='White'
def context():
 return dict(hennepin_packet=False,order_people=[dict(caption='Robin Original; Jamie Spouse',requested_caption='Robin Newname; Jamie Newname',primary=True,spouse=True,children='')],users=[Person('Robin Original')],district_court_county='Ramsey',judicial_district='Second',docket_number='',packet_caption='Robin Original; Jamie Spouse',packet_requested_caption='Robin Newname; Jamie Newname',request_for_applicant_name_change=True,applicant_requested_name='Robin Newname',request_for_applicant_name_change_on_birth_record=False,birth_requested_name='',request_sex_change_on_birth_record=False,replacement_birth_record=False,omit_former_name=False,omit_former_sex=False,spouse_included=True,spouse_requested_name='Jamie Newname',listed_spouse_summary='Jamie Spouse; February 3, 1972',listed_children_summary='None',included_children=[],included_children_names='',children_requested_names='',children_requested_summary='',has_felony=False,felony_summary='',land_summary='None',users1_other_statements='',release_people=[Person('Robin Original'),Person('Jamie Spouse')],felony_notice_data=[dict(current_name='Robin Original',requested_name='Robin Newname',kind='other_state',authority='Dane County, Wisconsin')],court_administrator_address='15 West Kellogg Boulevard',court_administrator_city_state_zip='St. Paul, MN 55102',inmate_name_change_reason='Match my identity',nonapplicant_parent_name='Other Parent',parent_last_location='Last known city',parent_last_employment='Unknown',parent_relatives='Relative at sample address',parent_search_efforts='Called the relative; no current address found',fee_legal_services=False,fee_public_assistance='none',fee_financial_details=True,fee_full_details=True,fee_household_size=1,fee_household_members='None',fee_income_sources='Job',fee_monthly_income=3000,fee_income_is_average=False,not_married=True,fee_other_household_income='None',fee_yearly_income=36000,fee_below_poverty_guideline=False,fee_expenses='Rent $2000; all other categories $0',fee_debt=1000,fee_cash=5,fee_accounts=10,fee_assets='Vehicle: none; home: none; other property: none',fee_other_reasons='Medical emergency',request_fee_waiver=True,users1_divorced_seeking_change_to_legal_name_on_birth_certificate=False,users1_inmate_in_correctional_facility_and_submitting_inmate_aff=False,parental_notice_needed=False)
def render(name,c):
 d=DocxTemplate(T/name)
 # Reject undeclared dependencies instead of allowing missing labels to render blank.
 from jinja2 import Environment,StrictUndefined
 with TemporaryDirectory() as tmp:
  d.render(c,jinja_env=Environment(undefined=StrictUndefined),autoescape=True)
  out=Path(tmp)/name;d.save(out);doc=Document(out)
  text='\n'.join(p.text for p in doc.paragraphs)
  text+='\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
  return text,doc
class PacketTemplates(unittest.TestCase):
 def test_core_release_and_order(self):
  c=context();text,_=render('criminal_history_releases.docx',c)
  self.assertEqual(text.count('Criminal History Check Release'),2)
  self.assertIn('Robin Original',text);self.assertIn('Jamie Spouse',text)
  text,_=render('proposed_name_change_order.docx',c)
  before,after=text.split('IT IS ORDERED that:',1)
  self.assertIn('Robin Newname',before);self.assertNotIn('Robin Newname',after)
  self.assertNotIn('Jamie Newname',after)
  self.assertIn('Judge of District Court',after)
 def test_all_conditional_forms_have_resolvable_labels(self):
  c=context()
  for name in ['felony_name_change_notices.docx','inmate_name_change_affidavit.docx','parental_hearing_notice.docx','parental_personal_service.docx','parental_publication_request.docx','Application_for_name_change_next_steps.docx']:
   with self.subTest(name=name):
    text,_=render(name,c);self.assertNotIn('{{',text)
 def test_fee_waiver_skip_routes_exclude_stale_financial_answers(self):
  c=context()
  text,_=render('fee_waiver_affidavit.docx',c)
  self.assertIn('Medical emergency',text);self.assertIn('36000.00',text)
  c.update(fee_legal_services=True,fee_financial_details=False,fee_full_details=False,fee_lawyer_name='Counsel Example',fee_lawyer_program='Legal Aid')
  text,_=render('fee_waiver_affidavit.docx',c)
  self.assertIn('Counsel Example',text);self.assertNotIn('Medical emergency',text);self.assertNotIn('36000.00',text)
  c.update(fee_legal_services=False,fee_public_assistance='listed',fee_benefits=NS(true_values=lambda:['SNAP']))
  text,_=render('fee_waiver_affidavit.docx',c)
  self.assertIn('SNAP',text);self.assertNotIn('36000.00',text)
 def test_hennepin_separate_orders(self):
  c=context();c['order_people']=[dict(caption='Robin Original',requested_caption='Robin Newname',primary=True,spouse=False,children=''),dict(caption='Jamie Spouse',requested_caption='Jamie Newname',primary=False,spouse=True,children='')]
  text,_=render('proposed_name_change_order.docx',c)
  self.assertEqual(text.count('IT IS ORDERED that:'),2)
  self.assertIn('Robin Newname',text);self.assertIn('Jamie Newname',text)
 def test_parent_service_title_is_separate_from_children(self):
  c=context();c.update(included_children_names='Avery Original; Casey Original',children_requested_names='Avery Newname; Casey Newname')
  _,doc=render('parental_personal_service.docx',c)
  self.assertTrue(any(p.text == 'Affidavit of Personal Service' for p in doc.paragraphs))
  self.assertTrue(any(p.text == 'Avery Original; Casey Original' for p in doc.paragraphs))
 def test_source_judicial_order_section_is_unmodified(self):
  original=Document(ROOT/'docassemble/ApplicationForNameChange/data/sources/court-originals/NAM107.docx')
  labeled=Document(T/'proposed_name_change_order.docx')
  start=next(i for i,p in enumerate(original.paragraphs) if 'IT IS ORDERED' in p.text)
  self.assertEqual([p._p.xml for p in original.paragraphs[start:]],[p._p.xml for p in labeled.paragraphs[start+1:start+1+len(original.paragraphs)-start]])
if __name__=='__main__':unittest.main()
