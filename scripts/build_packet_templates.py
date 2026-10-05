"""Label published court Word forms; preserve uncompleted judicial/execution sections.
Requires the archived originals and LibreOffice for legacy .doc conversion.
"""
from pathlib import Path
from copy import deepcopy
import subprocess
from tempfile import TemporaryDirectory
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'docassemble/ApplicationForNameChange/data'
OUT=DATA/'templates'; ORIG=ROOT/'reference/court-originals'

def replace(d,index,text):
 p=d.paragraphs[index]; p.clear();p.add_run(text)
 # The original uses a checkbox as a Word list bullet. The labeled text supplies
 # the checked/unchecked glyph, so remove only that duplicate list definition.
 if '☒' in text:
  for node in p._p.xpath('./w:pPr/w:numPr'):node.getparent().remove(node)

def loop(d,variable,collection):
 body=d.element.body; first=body[0];start=OxmlElement('w:p');first.addprevious(start)
 from docx.text.paragraph import Paragraph
 Paragraph(start,body).add_run('{%p for '+variable+' in '+collection+' %}')
 d.add_paragraph('{%p if not loop.last %}');d.add_page_break();d.add_paragraph('{%p endif %}');d.add_paragraph('{%p endfor %}')

def caption_table(d):
 t=d.tables[0]
 t.cell(0,0).text='State of Minnesota';t.cell(0,1).text=''
 t.cell(1,0).text='County: {{ district_court_county }}';t.cell(1,1).text=''
 t.cell(1,3).text='{{ judicial_district }}';t.cell(2,3).text='{{ docket_number }}'

with TemporaryDirectory() as tmp:
 for n in ['NAM103','NAM105','NAM205','NAM207']:
  subprocess.run(['libreoffice','--headless','-env:UserInstallation=file://'+tmp+'/profile','--convert-to','docx','--outdir',tmp,str(ORIG/(n+'.doc'))],check=True,capture_output=True)
 d=Document(Path(tmp)/'NAM103.docx')
 replace(d,10,'Full Name: {{ person.name.full() }}')
 replace(d,16,'{{ person.release_other_names }}');replace(d,18,'')
 replace(d,20,'Date of birth: {{ person.birthdate.format() }}')
 replace(d,21,'{{ "☒" if person.release_sex == "Female" else "☐" }} Female    {{ "☒" if person.release_sex == "Male" else "☐" }} Male    Race: {{ person.release_race }}')
 loop(d,'person','release_people');d.save(OUT/'criminal_history_releases.docx')
 d=Document(ORIG/'NAM107.docx');caption_table(d)
 for index,text in {
 3:"{{ order['caption'] }}",4:'',9:"{{ order['requested_caption'] }}",10:'',22:'of the application, and now live at: {{ users[0].address.on_one_line() }}',
 24:'in {{ users[0].address.county }} County.',26:'Name of applicant and date of birth: {{ users[0].name.full() }}; {{ users[0].birthdate.format() }}',
 28:'Name of spouse and date of birth: {{ listed_spouse_summary }}',29:'This application {{ "☒" if spouse_included else "☐" }} does  {{ "☐" if spouse_included else "☒" }} does not include spouse.',
 31:'Name(s) of minor children and date(s) of birth: {{ listed_children_summary }}',32:'',
 33:'{{ "☐" if included_children else "☒" }} This application does not include minor children listed above.',
 34:'{{ "☒" if included_children else "☐" }} This application includes the following minor children listed above: {{ included_children_names }}',
 37:'{{ "☒" if order[\'primary\'] and request_for_applicant_name_change else "☐" }} To have his/her name changed to {{ applicant_requested_name if order[\'primary\'] else "" }}',
 38:'{{ "☒" if order[\'primary\'] and request_for_applicant_name_change_on_birth_record else "☐" }} To have his/her name changed on birth records created or maintained by the Minnesota Department of Health to {{ birth_requested_name if order[\'primary\'] else "" }}',
 39:'{{ "☒" if order[\'primary\'] and request_sex_change_on_birth_record else "☐" }} To have his/her sex changed on birth records created or maintained by the Minnesota Department of Health to {{ sex_change if order[\'primary\'] and request_sex_change_on_birth_record else "" }}.',
 40:'{{ "☒" if order[\'primary\'] and replacement_birth_record else "☐" }} To have the Minnesota Department of Health issue and register a replacement birth record. Applicant further requests the prior birth record be kept confidential and the replacement birth record not to include any reference to Applicant’s {{ "☒" if order[\'primary\'] and omit_former_name else "☐" }} former name {{ "☒" if order[\'primary\'] and omit_former_sex else "☐" }} former sex.',
 41:'{{ "☒" if order[\'spouse\'] else "☐" }} To have the name of his/her spouse changed to {{ spouse_requested_name if order[\'spouse\'] else "" }}',
 42:'{{ "☒" if order[\'children\'] else "☐" }} To have the name(s) of his/her child (ren) changed to {{ order[\'children\'] }}',
 45:'{{ "☐" if has_felony else "☒" }} Has not been convicted of a felony in any state.',46:'{{ "☒" if has_felony else "☐" }} Has been convicted of a felony as follows: {{ felony_summary }}',47:'',
 55:'{{ land_summary }}',57:'Other: {{ users1_other_statements }}',58:''}.items():replace(d,index,text)
 # Never automate service/no-objection findings or anything after IT IS ORDERED.
 loop(d,'order','order_people');d.save(OUT/'proposed_name_change_order.docx')
 d=Document(ORIG/'NAM104.docx')
 for index,text in {
 1:'County of: {{ district_court_county }}    Court File Number: {{ docket_number }}',2:'Judicial District: {{ judicial_district }}    Case Type: Name Change',
 5:'{{ notice[\'current_name\'] }}',7:'{{ notice[\'requested_name\'] }}',10:'☒ Prosecuting Authority for:',
 11:'{{ "☐" if notice[\'kind\'] == "federal" else "☒" }} County and State: {{ notice[\'authority\'] if notice[\'kind\'] != "federal" else "" }}',
 12:'{{ "☒" if notice[\'kind\'] == "federal" else "☐" }} Federal District: {{ notice[\'authority\'] if notice[\'kind\'] == "federal" else "" }}',
 13:'{{ "☒" if notice[\'kind\'] != "minnesota" else "☐" }} Minnesota Attorney General (MN Attorney General must also be served if conviction is federal or out of state)',
 15:'This notice is to inform you that the Applicant has applied for a change of name in the {{ district_court_county }} County District Court by filing: (Check one)',
 16:'☒ Application for Name Change under Minn. Stat. § 259.10.',20:'Address: {{ court_administrator_address }}',21:'City/State/Zip: {{ court_administrator_city_state_zip }}',
 26:'Name: {{ users[0].name.full() }}',27:'County and state where signed: ____________________    Address: {{ users[0].address.address }}',
 28:'City/State/Zip: {{ users[0].address.line_two() }}',29:'Phone: {{ users[0].phone_number }}',30:'Email: {{ users[0].email }}'}.items():replace(d,index,text)
 loop(d,'notice','felony_notice_data');d.save(OUT/'felony_name_change_notices.docx')
 d=Document(Path(tmp)/'NAM105.docx');caption_table(d)
 for index,text in {4:'{{ users[0].name.full() }}',9:'{{ applicant_requested_name }}',16:'I, {{ users[0].name.full() }}, the applicant in this matter, make the following statement:',17:'☒ I am currently an inmate confined in a correctional facility, as defined in section 241.021, subdivision 1(f).',19:'☒ I have not at any time during my confinement requested a name change under section 259.10, other than this request.',21:'☒ The reason I am seeking a name change is: {{ inmate_name_change_reason }}',27:'☒ I request the court to issue its Order Granting Name Change.',33:'Name: {{ users[0].name.full() }}',34:'Address: {{ users[0].address.address }}',35:'City/State/Zip: {{ users[0].address.line_two() }}',36:'Telephone: {{ users[0].phone_number }}',37:'E-mail address: {{ users[0].email }}'}.items():replace(d,index,text)
 d.save(OUT/'inmate_name_change_affidavit.docx')
 d=Document(Path(tmp)/'NAM205.docx');caption_table(d)
 for index,text in {4:'{{ users[0].name.full() }}',9:'{{ included_children_names }}',14:'{{ children_requested_names }}',24:'18 years of age having been born on ______________ and that on ______________, I served the Application for a Name Change of a Minor and a notice of hearing upon {{ nonapplicant_parent_name }} at ______________________________ (address where documents were served) by handing a true and correct copy of the documents to him/her.'}.items():replace(d,index,text)
 # The court original shares the title line with a tabbed child-name blank.
 # Long names need their own paragraph so the title cannot be mistaken for a name.
 if 'Heading 1' not in d.styles:
  heading=d.styles.add_style('Heading 1',WD_STYLE_TYPE.PARAGRAPH)
  heading.font.name='Times New Roman';heading.font.size=Pt(14);heading.font.bold=True
  level=OxmlElement('w:outlineLvl');level.set(qn('w:val'),'0');heading.element.get_or_add_pPr().append(level)
 d.paragraphs[2].insert_paragraph_before('Affidavit of Personal Service',style='Heading 1')
 d.save(OUT/'parental_personal_service.docx')
 d=Document(Path(tmp)/'NAM207.docx');caption_table(d)
 for index,text in {2:'{{ users[0].name.full() }}',7:'On Behalf of: {{ included_children_names }} (Minor Name Change)',12:'{{ children_requested_names }}',22:'I have filed an Application for Name Change in {{ district_court_county }} County District Court for a change of name for the minor child (ren) from {{ included_children_names }} to {{ children_requested_names }}.',26:'2c. ☒ I do not know the address of the non-applicant parent.',28:'3. The last known location of the non-applicant parent was: {{ parent_last_location }}',31:'The last known location of the non-applicant parent’s employment was: {{ parent_last_employment }}',34:'5. The names and addresses of the non-applicant parent’s parents, brothers or sisters, children, and other close relatives are: {{ parent_relatives }}',37:'{{ parent_search_efforts }}'}.items():replace(d,index,text)
 d.save(OUT/'parental_publication_request.docx')
# A hearing notice is a supplement; dates/place can be completed once the court schedules it.
d=Document();d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(11)
for text in ['State of Minnesota — District Court','County: {{ district_court_county }}    Judicial district: {{ judicial_district }}','Court file number: {{ docket_number }}','Applicant: {{ users[0].name.full() }}','Notice of hearing — minor name change','To: {{ nonapplicant_parent_name }}','Children’s current names: {{ included_children_names }}','Requested names: {{ children_requested_names }}','A hearing on the application for these name changes will be held as follows:','Date: ____________________    Time: ____________________','Court location / remote appearance instructions: ________________________________________','________________________________________________________________________________','A copy of the application is attached.','Follow the court’s instructions for delivering this notice and proving delivery.']:
 d.add_paragraph(text)
d.save(OUT/'parental_hearing_notice.docx')
# Some LibreOffice .doc conversions use a nonstandard core-properties relationship.
# docxtpl then adds a second relationship/ZIP part, which LibreOffice rejects.
# Canonicalize that relationship on the converted templates before any rendering.
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
for name in ['criminal_history_releases','inmate_name_change_affidavit','parental_personal_service','parental_publication_request']:
 path=OUT/(name+'.docx')
 with ZipFile(path) as z:
  parts={info.filename:z.read(info.filename) for info in z.infolist()}
 rels=etree.fromstring(parts['_rels/.rels'])
 for rel in rels:
  if rel.get('Target')=='docProps/core.xml':rel.set('Type','http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties')
 parts['_rels/.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
 with ZipFile(path,'w',ZIP_DEFLATED) as z:
  for key,value in parts.items():z.writestr(key,value)
# Compact conversion-only empty spacing paragraphs on one-page forms. Preserve
# the signature blank and page breaks, and allow long answers to expand naturally.
for name in ['criminal_history_releases','inmate_name_change_affidavit']:
 path=OUT/(name+'.docx');doc=Document(path)
 for para in list(doc.paragraphs):
  if not para.text.strip() and not para._p.xpath('.//w:br') and not para._p.xpath('.//w:sectPr'):
   para._p.getparent().remove(para._p)
 for para in doc.paragraphs:
  para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(4)
  if 'Signature' in para.text or 'Your Signature' in para.text:
   para.paragraph_format.space_before=Pt(14)
 doc.save(path)
